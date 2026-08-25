"""
Stage 4. Turns the derived cost data into a resource-level work breakdown structure:
every item resolved into the individual people, plant and materials needed to build it,
each priced from a single editable rate library.

    python3 build_wbs.py [work_dir] [output.xlsx]

The model separates the two things that behave differently:

  * CONSTANTS — the crew, the plant spread, the materials and the quantities of each.
    This is "how the work gets done". It does not change when the market does.
  * RATES — what an hour of a ganger or a day of a backacter costs. This is what you
    re-rate to suit your own business.

Requires openpyxl. work_dir must hold out.json / res.json from derive_costs.py.
"""
import json, os, sys, re, collections, statistics
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
W = (sys.argv[1] if len(sys.argv) > 1 else HERE).rstrip("/") + "/"
OUT = sys.argv[2] if len(sys.argv) > 2 else W + "InfraBid_WBS_Rate_Model.xlsx"

out = json.load(open(W + "out.json"))
res = json.load(open(W + "res.json"))

# optional 3rd arg: build a small sample instead of the full book, so the formula
# grammar can be verified by a real recalculation before generating 40k+ formulas.
SAMPLE = int(sys.argv[3]) if len(sys.argv) > 3 else 0
if SAMPLE:
    keep = {o['src'] for o in out[:SAMPLE]}
    out = [o for o in out if o['src'] in keep]

# ---------------------------------------------------------------- Irish basis
FX, LABOUR_F, PLANT_F, MATERIAL_F = 1.1694, 1.3415, 1.00, 1.00
SC_MIX = (0.35, 0.30, 0.35)
SC_F = round(SC_MIX[0]*LABOUR_F + SC_MIX[1]*PLANT_F + SC_MIX[2]*MATERIAL_F, 4)

# ============================================================ parse resources
TOTAL = re.compile(r'^total\s+(gang\s+)?rate\s*/\s*hour$', re.I)
COUNT = re.compile(r'^\s*(\d+)\s+(.*)$')
WORDN = {'one':1,'two':2,'three':3,'four':4,'five':5,'six':6,'seven':7,'eight':8}
# the source writes part-time allocation two ways: '- 33% of time' and '(33% of time)',
# sometimes '(50% time)'. Consume the brackets too, or an empty '()' is left behind.
PCT   = re.compile(r'\s*[\(\[]?\s*[-–;,]?\s*(\d+(?:\.\d+)?)\s*%\s*(?:of\s*)?time\s*[\)\]]?', re.I)
UNITWORD = re.compile(r'^(kg|kgs|t|te|tonne|tonnes|mm|cm|m|m2|m3|ft|in|hp|kw|litre|l|cwt|deg|"|\u2033)\b', re.I)

def split_line(txt, digit_counts):
    """'2 unskilled operatives - 33% of time' -> (2, 0.33, 'unskilled operatives')

    digit_counts=False for plant: a leading number there is a SIZE, not a quantity
    ('1000 kg hydraulic breaker', '8 tonne wheeled backacter', '3.7 m3/min compressor'),
    and reading it as a count inflates the machine count a thousandfold."""
    t = txt.strip(); frac = 1.0
    m = PCT.search(t)
    if m:
        frac = float(m.group(1))/100.0
        t = PCT.sub("", t).strip(" -–;,")
    t = re.sub(r'\(\s*\)', '', t).strip(" -–;,")
    n = 1
    w = t.split(" ", 1)
    if w and w[0].lower() in WORDN:
        n = WORDN[w[0].lower()]; t = w[1].strip() if len(w) > 1 else t
    elif digit_counts:
        m = COUNT.match(t)
        if m and 1 <= int(m.group(1)) <= 12 and not UNITWORD.match(m.group(2)):
            n = int(m.group(1)); t = m.group(2).strip()
    return n, frac, t

def canon(s):
    t = re.sub(r'\s+', ' ', s.lower().strip())
    t = t.replace("generally", "general")
    t = re.sub(r'\bganger or chargehand\b', 'ganger/chargehand', t)
    t = re.sub(r'\b(operatives|operators|labourers|gangers|drivers|fitters|joiners)\b',
               lambda m: m.group(1)[:-1], t)
    return t.strip(" .,;-")

blocks = collections.OrderedDict()
for r in res:
    blocks.setdefault((r[1], r[2], r[3]), []).append(r)

gangs, spreads, raw_materials = {}, {}, []
for (cls, sec, sub), rows in blocks.items():
    if sec == 'RESOURCES - MATERIAL':
        for r in rows:
            if r[8] is not None and r[8] > 0:
                raw_materials.append(dict(cls=cls, group=sub, desc=r[4],
                                          unit=(r[9] or "nr"), rate=r[8], src=r[12]))
        continue
    kind = 'labour' if 'LABOUR' in sec.upper() else 'plant'
    col = 6 if kind == 'labour' else 7
    lines = []
    for r in rows:
        if TOTAL.match(str(r[4]).strip()): continue
        v = r[col] if r[col] is not None else (r[6] if r[6] is not None else r[7])
        if v is None or v <= 0: continue
        n, frac, name = split_line(r[4], digit_counts=(kind == 'labour'))
        lines.append(dict(n=n, frac=frac, name=canon(name), unit_rate=v/(n*frac)))
    if lines:
        (gangs if kind == 'labour' else spreads)[(cls, sub)] = lines

# ---------------------------------------------------------------- library
def build_pool(defs, prefix, kind):
    agg = collections.defaultdict(list)
    for lines in defs.values():
        for l in lines: agg[l['name']].append(l['unit_rate'])
    pool = {}
    for i, (name, rates) in enumerate(sorted(agg.items()), 1):
        pool[name] = dict(code=f"{prefix}{i:03d}", kind=kind, desc=name,
                          unit="hr", rate=round(statistics.median(rates), 4),
                          n_uses=len(rates))
    return pool

LAB = build_pool(gangs,   "L", "Labour")
PLT = build_pool(spreads, "P", "Plant")

# materials: the source's own price list, plus the pure-supply items it prices as rates
MAT = {}
def add_mat(desc, unit, rate, group):
    key = canon(desc) + "|" + (unit or "nr")
    if key in MAT:
        MAT[key]['rates'].append(rate); return MAT[key]
    MAT[key] = dict(code=f"M{len(MAT)+1:03d}", kind="Material", desc=desc.strip(),
                    unit=unit or "nr", rates=[rate], group=group)
    return MAT[key]

for m in raw_materials:
    add_mat(m['desc'], m['unit'], m['rate'], m['group'])
# Class F concrete and the other pure-material items are themselves a material price list
for o in out:
    if o['M'] > 0 and o['L'] == 0 and o['P'] == 0 and o['SC'] == 0 and o['unit']:
        add_mat(f"{o['sub']}; {o['item']}" if o['sub'] and o['sub'] != o['item'] else o['item'],
                o['unit'], o['M'], o['cls'])
for v in MAT.values():
    v['rate'] = round(statistics.median(v['rates']), 4)

# ---------------------------------------------------------------- crew rates
def crew_rate(lines, pool):
    return round(sum(pool[l['name']]['rate'] * l['n'] * l['frac'] for l in lines), 4)

GANG = {k: dict(cls=k[0], name=k[1], lines=v, rate=crew_rate(v, LAB)) for k, v in gangs.items()}
SPRD = {k: dict(cls=k[0], name=k[1], lines=v, rate=crew_rate(v, PLT)) for k, v in spreads.items()}

# ---------------------------------------------------------------- assignment
STOP = set("""the a an of to for and or in on with by at from as is are be per each no not
gang gangs rate rates build up build-up resources resource hour hourly gauge nr mm m2 m3
class general""".split())

def toks(s):
    return set(w for w in re.findall(r'[a-z]+', s.lower()) if w not in STOP and len(w) > 2)

def pick(pool_by_cls, cls, text, fallback_all):
    """Choose the crew/spread whose name best fits this item's section and sub-heading."""
    cands = pool_by_cls.get(cls) or fallback_all
    if not cands: return None, "none"
    if len(cands) == 1: return cands[0], "only crew in class"
    it = toks(text)
    best, score = None, -1
    for c in cands:
        s = len(it & toks(c['name']))
        if s > score: best, score = c, s
    if score <= 0:
        return cands[0], "class default"
    return best, "name match"

gang_by_cls = collections.defaultdict(list)
for g in GANG.values(): gang_by_cls[g['cls']].append(g)
sprd_by_cls = collections.defaultdict(list)
for s in SPRD.values(): sprd_by_cls[s['cls']].append(s)
ALL_G = sorted(GANG.values(), key=lambda g: -len(g['lines']))
ALL_P = sorted(SPRD.values(), key=lambda s: -len(s['lines']))

# ---------------------------------------------------------------- materials
def mat_candidates(cls):
    return [v for v in MAT.values() if v.get('group') == cls or True]

MAT_BY_CLS = collections.defaultdict(list)
for m in raw_materials:
    key = canon(m['desc']) + "|" + (m['unit'] or "nr")
    if key in MAT: MAT_BY_CLS[m['cls']].append(MAT[key])
for o in out:
    if o['M'] > 0 and o['L'] == 0 and o['P'] == 0 and o['SC'] == 0 and o['unit']:
        d = f"{o['sub']}; {o['item']}" if o['sub'] and o['sub'] != o['item'] else o['item']
        key = canon(d) + "|" + o['unit']
        if key in MAT: MAT_BY_CLS[o['cls']].append(MAT[key])

SIZE = re.compile(r'(\d+(?:\.\d+)?)\s*(?:mm|m2|m3|m\b|kg|t\b|tonne)?')
def sizes(s): return set(re.findall(r'\d+(?:\.\d+)?', s))

def match_material(o):
    """Find a named material in the same class that fits this item. Size must agree."""
    if o['M'] <= 0: return None, 0
    text = f"{o['sub']} {o['item']}"
    it, isz = toks(text), sizes(text)
    best, bs = None, 0
    for m in MAT_BY_CLS.get(o['cls'], []):
        if m['unit'] != o['unit']: continue
        mt, ms = toks(m['desc']), sizes(m['desc'])
        s = len(it & mt)
        if ms and isz:
            if ms & isz: s += 3
            else: continue                      # a different diameter is a different material
        if s > bs: best, bs = m, s
    if best is None or bs < 3: return None, bs
    # units are forced equal above, so the derived quantity should sit near 1 (plus wastage).
    # Anything far from that means the wrong material matched — fall back to item-specific.
    q = o['M'] / best['rate']
    if not (0.4 <= q <= 3.0): return None, bs
    return best, bs

# ============================================================ build the WBS
wbs, items, unmatched_mat = [], [], 0
stat = collections.Counter()

for o in out:
    ref = o['src']
    text = f"{o['sec']} {o['sub']} {o['item']}"
    lines = []

    # ---- labour
    if o['L'] > 0:
        g, how = pick(gang_by_cls, o['cls'], text, ALL_G)
        if g and g['rate'] > 0:
            hrs = o['L'] * LABOUR_F * FX / (g['rate'] * LABOUR_F * FX)   # = L / crew rate, in £
            hrs = o['L'] / g['rate']
            stat[f"labour: {how}"] += 1
            for l in g['lines']:
                q = round(hrs * l["n"] * l["frac"], 6)
                if q <= 0: continue
                lines.append(("Labour", LAB[l['name']]['code'], LAB[l['name']]['desc'],
                              q, "hr", g['name']))
        else:
            stat["labour: no crew available"] += 1
            lines.append(("Labour", "L000", "Composite gang (no crew published)",
                          1, "item", "—"))
    # ---- plant
    if o['P'] > 0:
        s, how = pick(sprd_by_cls, o['cls'], text, ALL_P)
        if s and s['rate'] > 0:
            hrs = o['P'] / s['rate']
            stat[f"plant: {how}"] += 1
            for l in s['lines']:
                q = round(hrs * l["n"] * l["frac"], 6)
                if q <= 0: continue
                lines.append(("Plant", PLT[l['name']]['code'], PLT[l['name']]['desc'],
                              q, "hr", s['name']))
        else:
            stat["plant: no spread available"] += 1
            lines.append(("Plant", "P000", "Composite plant spread (none published)",
                          1, "item", "—"))
    # ---- material
    if o['M'] > 0:
        m, sc = match_material(o)
        if m:
            stat["material: matched to library"] += 1
            lines.append(("Material", m['code'], m['desc'],
                          round(o["M"]/m["rate"], 6), m['unit'], m['group']))
        else:
            unmatched_mat += 1
            stat["material: item-specific"] += 1
            lines.append(("Material", f"MX{ref}",
                          (o['sub'] if o['sub'] and o['sub'] != o['item'] else o['item'])[:120],
                          1, o['unit'] or "item", "item-specific"))
    # ---- sub-contract
    if o['SC'] > 0:
        stat["sub-contract package"] += 1
        lines.append(("Sub-contract", f"SX{ref}",
                      f"{o['sub'] or o['item']}"[:120], 1, o['unit'] or "item",
                      "specialist package"))

    for i, (kind, code, desc, qty, unit, srcname) in enumerate(lines, 1):
        wbs.append(dict(ref=ref, cls=o['cls'], sec=o['sec'], sub=o['sub'], item=o['item'],
                        iunit=o['unit'], line=i, kind=kind, code=code, desc=desc,
                        qty=qty, unit=unit, crew=srcname))
    items.append(o)

# item-specific resources become library entries too, so every line has a rate
EXTRA = {}
for w in wbs:
    if w['code'].startswith(("MX", "SX")):
        o = next(x for x in out if x['src'] == w['ref'])
        EXTRA[w['code']] = dict(code=w['code'],
                                kind=w['kind'], desc=w['desc'], unit=w['unit'],
                                rate=(o['M'] if w['kind'] == "Material" else o['SC']),
                                group="item-specific")

print(f"WBS lines: {len(wbs):,}  over {len(items):,} items "
      f"(avg {len(wbs)/len(items):.1f} lines/item)")
print(f"library: labour {len(LAB)}  plant {len(PLT)}  material {len(MAT)}  "
      f"item-specific {len(EXTRA)}")
print("\nassignment:")
for k, v in stat.most_common(): print(f"   {v:6,d}  {k}")

json.dump(dict(wbs=wbs, LAB=LAB, PLT=PLT,
               MAT={k: {kk: vv for kk, vv in v.items() if kk != 'rates'} for k, v in MAT.items()},
               EXTRA=EXTRA,
               GANG={f"{k[0]}||{k[1]}": dict(cls=v['cls'], name=v['name'], rate=v['rate'],
                                             lines=v['lines']) for k, v in GANG.items()},
               SPRD={f"{k[0]}||{k[1]}": dict(cls=v['cls'], name=v['name'], rate=v['rate'],
                                             lines=v['lines']) for k, v in SPRD.items()}),
          open(W + "wbs.json", "w"))
print("\nwrote", W + "wbs.json")

# ============================================================ workbook
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
NAVY, INK, BORDERC = "133198", "15234E", "E2EAF6"
EUR   = '€#,##0.00;-€#,##0.00;"–"'
EUR4  = '€#,##0.0000'
GBP   = '£#,##0.00;-£#,##0.00;"–"'
QTY   = '#,##0.000;-#,##0.000;"–"'
HDRF  = PatternFill("solid", fgColor=NAVY)
INPUTF= PatternFill("solid", fgColor="FFF2CC")     # yellow = you edit this
CALCF = PatternFill("solid", fgColor="EAF1FA")
BANDF = PatternFill("solid", fgColor="F4F8FD")
THIN  = Side(style="thin", color=BORDERC)
BOX   = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TITLE = Font(name=FONT, sz=15, bold=True, color=NAVY)
SUBF  = Font(name=FONT, sz=10, italic=True, color="595959")
BODY  = Font(name=FONT, sz=10)
BOLDF = Font(name=FONT, sz=10, bold=True)
HDRT  = Font(name=FONT, sz=10, bold=True, color="FFFFFF")

def head(ws, cols, title, sub, row=4):
    ws.cell(1, 1, title).font = TITLE
    ws.cell(2, 1, sub).font = SUBF
    for i, (h, w, fmt, al) in enumerate(cols, 1):
        c = ws.cell(row, i, h); c.fill = HDRF; c.font = HDRT; c.border = BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 28
    ws.freeze_panes = ws.cell(row+1, 1)

def put(ws, cols, data, start):
    for ri, r in enumerate(data, start):
        for ci, (v, (h, w, fmt, al)) in enumerate(zip(r, cols), 1):
            c = ws.cell(ri, ci, v); c.font = BODY; c.border = BOX
            if fmt: c.number_format = fmt
            c.alignment = Alignment(horizontal=al, vertical="top")

wb = Workbook(); wb.remove(wb.active)

# ---------------------------------------------------------------- 1. library
lib = []
for pool, scope in ((LAB, "Shared"), (PLT, "Shared"), (MAT, "Shared")):
    for v in pool.values():
        lib.append([v['code'], v['kind'], v['desc'], v['unit'], round(v['rate'], 4), scope,
                    v.get('group', "") or ""])
for v in EXTRA.values():
    lib.append([v['code'], v['kind'], v['desc'], v['unit'], round(v['rate'], 4),
                "Item-specific", v['group']])
order = {"Labour": 0, "Plant": 1, "Material": 2, "Sub-contract": 3}
lib.sort(key=lambda r: (0 if r[5] == "Shared" else 1, order.get(r[1], 9), r[0]))
use = collections.Counter(w['code'] for w in wbs)

ws = wb.create_sheet("Rate Library")
LCOLS = [("Code", 11, None, "center"), ("Type", 13, None, "left"),
         ("Resource", 62, None, "left"), ("Unit", 8, None, "center"),
         ("Base rate £", 13, GBP, "right"), ("Factor", 9, "0.0000", "right"),
         ("Derived rate €", 14, EUR, "right"), ("YOUR RATE €", 14, EUR, "right"),
         ("Rate used €", 14, EUR, "right"), ("Scope", 13, None, "left"),
         ("Used on", 10, "#,##0", "right"), ("Source group", 52, None, "left")]
head(ws, LCOLS,
     "Rate Library — every person, machine and material, priced once",
     "This is the tab you re-rate. Type your own figure in the yellow YOUR RATE € column and it "
     "overrides the derived one everywhere; leave it blank to keep the derived rate. Nothing else "
     "needs to change — the quantities in the WBS are how the work gets done, not what it costs.",
     row=10)
ws.cell(4, 1, "Conversion basis").font = BOLDF
for i, (lbl, val) in enumerate([("GBP to EUR", FX), ("Labour factor", LABOUR_F),
                                ("Plant factor", PLANT_F), ("Material factor", MATERIAL_F),
                                ("Sub-contract factor", SC_F)]):
    ws.cell(4+i, 2, lbl).font = BODY
    c = ws.cell(4+i, 3, val); c.number_format = "0.0000"; c.font = BOLDF
    c.fill = INPUTF; c.border = BOX; c.alignment = Alignment(horizontal="right")
FXC = "$C$4"
FMAP = {"Labour": "$C$5", "Plant": "$C$6", "Material": "$C$7", "Sub-contract": "$C$8"}

rows = [[r[0], r[1], r[2], r[3], r[4], None, None, None, None, r[5], use.get(r[0], 0), r[6]]
        for r in lib]
put(ws, LCOLS, rows, 11)
for i, r in enumerate(lib):
    n = 11 + i
    ws.cell(n, 6, f"={FMAP.get(r[1], '$C$8')}").number_format = "0.0000"
    ws.cell(n, 7, f"=ROUND(E{n}*F{n}*{FXC},4)").number_format = EUR4
    ws.cell(n, 9, f"=IF(H{n}=\"\",G{n},H{n})").number_format = EUR4
    for col in (6, 7, 9):
        ws.cell(n, col).font = BODY; ws.cell(n, col).border = BOX
        ws.cell(n, col).alignment = Alignment(horizontal="right")
    ws.cell(n, 8).fill = INPUTF
    ws.cell(n, 9).fill = CALCF; ws.cell(n, 9).font = BOLDF
LIB_LAST = 10 + len(lib)
ws.auto_filter.ref = f"A10:L{LIB_LAST}"

# ---------------------------------------------------------------- 2. WBS
ws2 = wb.create_sheet("WBS")
WCOLS = [("Item ref", 10, "0", "center"), ("Class", 30, None, "left"),
         ("Section", 30, None, "left"), ("Sub-heading", 40, None, "left"),
         ("Item", 46, None, "left"), ("Item unit", 9, None, "center"),
         ("Line", 6, "0", "center"), ("Type", 13, None, "left"),
         ("Code", 11, None, "center"), ("Resource", 56, None, "left"),
         ("Qty", 12, QTY, "right"), ("Qty unit", 9, None, "center"),
         ("Rate €", 13, EUR, "right"), ("Line cost €", 14, EUR, "right"),
         ("Crew / spread", 46, None, "left")]
head(ws2, WCOLS,
     "WBS — what it takes to build each item",
     "One row per resource. Quantities are the constants: the crew, the plant spread and the "
     "materials. Rates come from the Rate Library, so re-rating there flows straight through here.")
put(ws2, WCOLS, [[w['ref'], w['cls'], w['sec'], w['sub'], w['item'], w['iunit'], w['line'],
                  w['kind'], w['code'], w['desc'], w['qty'], w['unit'], None, None, w['crew']]
                 for w in wbs], 5)
for i in range(len(wbs)):
    n = 5 + i
    ws2.cell(n, 13, f"=INDEX('Rate Library'!$I$11:$I${LIB_LAST},"
                    f"MATCH(I{n},'Rate Library'!$A$11:$A${LIB_LAST},0))").number_format = EUR4
    ws2.cell(n, 14, f"=ROUND(K{n}*M{n},2)").number_format = EUR
    for col in (13, 14):
        ws2.cell(n, col).font = BODY; ws2.cell(n, col).border = BOX
        ws2.cell(n, col).alignment = Alignment(horizontal="right")
    ws2.cell(n, 14).font = BOLDF
WBS_LAST = 4 + len(wbs)
ws2.auto_filter.ref = f"A4:O{WBS_LAST}"

# ---------------------------------------------------------------- 3. item cost
ws3 = wb.create_sheet("Item Cost")
ICOLS = [("Item ref", 10, "0", "center"), ("Class", 30, None, "left"),
         ("Section", 30, None, "left"), ("Sub-heading", 42, None, "left"),
         ("Item", 50, None, "left"), ("Unit", 9, None, "center"),
         ("Labour €", 13, EUR, "right"), ("Plant €", 13, EUR, "right"),
         ("Material €", 13, EUR, "right"), ("Sub-contract €", 14, EUR, "right"),
         ("TOTAL COST €", 15, EUR, "right"), ("WBS lines", 10, "0", "right"),
         ("Book-derived €", 14, EUR, "right"), ("Variance €", 12, EUR, "right")]
head(ws3, ICOLS,
     "Item Cost — built up from the WBS",
     "Every figure here is a live sum of the WBS lines for that item. Change a rate in the Rate "
     "Library and these move. 'Book-derived €' is the cost carried over from the price book, for "
     "comparison only.")
nlines = collections.Counter(w['ref'] for w in wbs)
put(ws3, ICOLS, [[o['src'], o['cls'], o['sec'], o['sub'], o['item'], o['unit'],
                  None, None, None, None, None, nlines[o['src']],
                  round((o['L']*LABOUR_F + o['P']*PLANT_F + o['M']*MATERIAL_F
                         + o['SC']*SC_F)*FX, 2), None] for o in items], 5)
W_ = f"WBS!$N$5:$N${WBS_LAST}"; R_ = f"WBS!$A$5:$A${WBS_LAST}"; T_ = f"WBS!$H$5:$H${WBS_LAST}"
for i in range(len(items)):
    n = 5 + i
    for col, kind in ((7, "Labour"), (8, "Plant"), (9, "Material"), (10, "Sub-contract")):
        ws3.cell(n, col, f'=SUMIFS({W_},{R_},$A{n},{T_},"{kind}")').number_format = EUR
    ws3.cell(n, 11, f"=ROUND(SUM(G{n}:J{n}),2)").number_format = EUR
    ws3.cell(n, 14, f"=ROUND(K{n}-M{n},2)").number_format = EUR
    for col in (7, 8, 9, 10, 11, 14):
        ws3.cell(n, col).font = BODY; ws3.cell(n, col).border = BOX
        ws3.cell(n, col).alignment = Alignment(horizontal="right")
    ws3.cell(n, 11).font = BOLDF; ws3.cell(n, 11).fill = CALCF
ITEM_LAST = 4 + len(items)
ws3.auto_filter.ref = f"A4:N{ITEM_LAST}"

# ---------------------------------------------------------------- 4. crews
ws4 = wb.create_sheet("Gangs and Spreads")
GCOLS = [("Kind", 11, None, "left"), ("Class", 30, None, "left"),
         ("Crew / spread", 52, None, "left"), ("Code", 11, None, "center"),
         ("Resource", 56, None, "left"), ("Number", 9, "0", "right"),
         ("Time share", 11, "0%", "right"), ("Effective units", 13, "0.000", "right"),
         ("Rate €/hr", 13, EUR, "right"), ("Cost €/hr", 13, EUR, "right")]
head(ws4, GCOLS, "Gangs and Spreads — the crews behind the WBS",
     "The composition published in the source. An item's labour hours are its labour cost divided "
     "by its crew's hourly cost, so the crew stays fixed and the hours flex.")
grows = []
for kind, pool, lib_pool in (("Labour gang", GANG, LAB), ("Plant spread", SPRD, PLT)):
    for g in sorted(pool.values(), key=lambda x: (x['cls'], x['name'])):
        for l in g['lines']:
            r = lib_pool[l['name']]
            grows.append([kind, g['cls'], g['name'], r['code'], r['desc'],
                          l['n'], l['frac'], round(l['n']*l['frac'], 3), None, None])
put(ws4, GCOLS, grows, 5)
for i, gr in enumerate(grows):
    n = 5 + i
    ws4.cell(n, 9, f"=INDEX('Rate Library'!$I$11:$I${LIB_LAST},"
                   f"MATCH(D{n},'Rate Library'!$A$11:$A${LIB_LAST},0))").number_format = EUR4
    ws4.cell(n, 10, f"=ROUND(H{n}*I{n},2)").number_format = EUR
    for col in (9, 10):
        ws4.cell(n, col).font = BODY; ws4.cell(n, col).border = BOX
        ws4.cell(n, col).alignment = Alignment(horizontal="right")
ws4.auto_filter.ref = f"A4:J{4+len(grows)}"

# ---------------------------------------------------------------- 5. read me
ws5 = wb.create_sheet("Read me")
for col, w in zip("ABCD", (4, 42, 22, 104)): ws5.column_dimensions[col].width = w
ws5.cell(1, 2, "WBS Rate Model — how it works").font = TITLE
ws5.cell(2, 2, "Spon's Civil Engineering rates, resolved into the people, plant and materials "
               "each item needs, on an Irish basis.").font = SUBF
r = 4
def h2(t):
    global r
    for c in range(2, 5):
        ws5.cell(r, c).fill = HDRF; ws5.cell(r, c).border = BOX
    ws5.cell(r, 2, t).font = Font(name=FONT, sz=11, bold=True, color="FFFFFF"); r += 1
def ln(a, b=None, c=None, bold=False):
    global r
    ws5.cell(r, 2, a).font = Font(name=FONT, sz=10, bold=bold)
    ws5.cell(r, 2).alignment = Alignment(vertical="top", wrap_text=True)
    if b is not None:
        cc = ws5.cell(r, 3, b); cc.font = Font(name=FONT, sz=10, bold=True)
        cc.alignment = Alignment(horizontal="right")
    if c:
        cc = ws5.cell(r, 4, c); cc.font = Font(name=FONT, sz=10)
        cc.alignment = Alignment(vertical="top", wrap_text=True)
    r += 1

h2("The idea")
ln("Constants and rates are separated.", None,
   "How the work gets done — the crew, the plant spread, the materials and the quantity of each — "
   "sits in the WBS and does not change when the market does. What an hour or a tonne costs sits in "
   "the Rate Library. Re-rate there and every item re-prices.")
ln("To use your own rates", None,
   "Open 'Rate Library', type your figure in the yellow YOUR RATE € column. It overrides the derived "
   "rate for that resource everywhere it is used. Leave it blank to keep the derived rate. You can "
   "also change the five factors at the top of that sheet to move a whole category at once.")
r += 1

h2("The sheets")
ln("Rate Library", f"{len(lib):,} rows",
   "Every person, machine and material, priced once. Shared resources first, then item-specific ones.")
ln("WBS", f"{len(wbs):,} rows",
   "One row per resource per item — the work breakdown. Qty x Rate = Line cost.")
ln("Item Cost", f"{len(items):,} rows",
   "One row per item, summing its WBS lines into labour, plant, material and sub-contract.")
ln("Gangs and Spreads", f"{len(grows):,} rows",
   "The crew and plant-spread compositions published in the source, and what each costs per hour.")
r += 1

h2("Where the numbers come from")
ln("Crews and spreads", f"{len(GANG)+len(SPRD)}",
   "Published in the price book as gang and plant build-ups, one set per CESMM class. Parsed into "
   "individual roles and machines with their hourly rates.")
ln("Labour resources", f"{len(LAB)}",
   "Each grade's hourly rate cross-checks exactly everywhere it appears in the book, so these are solid.")
ln("Plant resources", f"{len(PLT)}", "Individual machines from the plant build-ups.")
ln("Materials, shared", f"{len(MAT)}",
   "The book's own material price list, plus the items it prices as pure supply (concrete mixes, rail, "
   "pipes). Roughly 69% of all material value links to one of these.")
ln("Materials, item-specific", f"{len(EXTRA):,}",
   "Where the book gives a material value but never names the material, the line carries the item's own "
   "description. Still fully re-rateable, just not shared between items.")
r += 1

h2("How the hours were set")
ln("Hours flex, crews do not.", None,
   "An item's labour hours are its labour cost divided by its crew's hourly cost; plant hours likewise. "
   "So every item reconciles to the price book exactly while naming a real crew.")
ln("Why not the book's own gang hours?", None,
   "The book states a 'Total Gang Rate/Hour' for each crew that does not equal the sum of that crew's "
   "own lines — the median gap is 1.4% and the worst is far larger. The individual rates cross-check "
   "perfectly across every gang they appear in, so the line rates were trusted and the totals were not.")
r += 1

h2("What to check before using this")
for t in ["Crew selection is inferred. Where the class publishes more than one gang, the crew is chosen "
          "by matching the gang's name against the item's section and sub-heading; failing that, the "
          "class's first gang is used. The cost is right either way, but the named crew may not be.",
          "Material quantities are derived, not quoted. Where a material is matched, quantity is the "
          "item's material value divided by that material's unit rate — accepted only when it lands "
          "between 0.4 and 3.0 units, which is where a genuine match sits.",
          "653 items are specialist sub-contract packages with no build-up in the source. They carry a "
          "single sub-contract line rather than an invented crew.",
          "Rates are net cost: no overhead, profit, preliminaries or VAT.",
          "The source is an OCR extraction of a scanned book and some figures are mis-read."]:
    ln("•", None, t)

for s in wb.worksheets: s.sheet_view.showGridLines = False
wb.move_sheet("Read me", offset=-(len(wb.worksheets)-1))
wb.save(OUT)
print(f"\nsaved {OUT}")
print(f"  Rate Library {len(lib):,} | WBS {len(wbs):,} | Item Cost {len(items):,} | Crews {len(grows):,}")
